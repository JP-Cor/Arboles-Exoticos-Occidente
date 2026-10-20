# Ejecutar desde la raíz del proyecto: python scripts/actualizar_inventario.py archivo.xlsx
# Requiere: pip install openpyxl
import json,sys,pathlib
from openpyxl import load_workbook
raiz=pathlib.Path(__file__).resolve().parents[1]
archivo=pathlib.Path(sys.argv[1])
w=load_workbook(archivo,read_only=True,data_only=True)
s=w['Catálogo completo'] if 'Catálogo completo' in w.sheetnames else w.active
keys=['id','categoria','especie','precio','tamano','propagacion','observaciones','disponibilidad','cientifico','ingles','validacion']
previo=raiz/'Inventario'/'inventario.json'
fotos={}
if previo.exists():
 for p in json.loads(previo.read_text(encoding='utf-8')):fotos[p['id']]=p.get('foto','')
registros=[]
for row in list(s.values)[1:]:
 if not isinstance(row[0],(int,float)):continue
 d={k:(row[i] if i<len(row) and row[i] is not None else '') for i,k in enumerate(keys)}
 d['id']=int(d['id']);d['foto']=fotos.get(d['id'],'');registros.append(d)
previo.write_text(json.dumps(registros,ensure_ascii=False,indent=2),encoding='utf-8')
print(f'Actualizadas {len(registros)} presentaciones en {previo}')
